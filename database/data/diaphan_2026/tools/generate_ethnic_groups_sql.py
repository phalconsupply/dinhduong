from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DOCX = ROOT / "dantoc" / "dantoc.docx"
OUTPUT_SQL = ROOT / "dantoc" / "mysql_ethnic_groups.sql"
OUTPUT_JSON = ROOT / "dantoc" / "ethnic_groups.json"


def clean_text(value: str) -> str:
    return " ".join(value.replace("\xa0", " ").split())


def sql_str(value: str | None) -> str:
    if value is None or value == "":
        return "NULL"
    return "'" + value.replace("\\", "\\\\").replace("'", "''") + "'"


def extract_groups() -> list[dict[str, str | int | bool]]:
    document = Document(str(SOURCE_DOCX))
    if not document.tables:
        raise RuntimeError(f"No table found in {SOURCE_DOCX}")

    table = document.tables[0]
    rows = table.rows[1:]
    groups: list[dict[str, str | int | bool]] = []

    for row in rows:
        cells = [clean_text(cell.text) for cell in row.cells]
        if len(cells) < 5 or not cells[0].isdigit():
            continue

        sort_order = int(cells[0])
        groups.append(
            {
                "id": sort_order,
                "code": f"{sort_order:02d}",
                "name": cells[1],
                "other_names": cells[2],
                "subgroups": cells[3],
                "residence_area": cells[4],
                "is_majority": sort_order == 1,
                "sort_order": sort_order,
            }
        )

    return groups


def write_json(groups: list[dict[str, str | int | bool]]) -> None:
    import json

    OUTPUT_JSON.write_text(
        json.dumps(groups, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def write_sql(groups: list[dict[str, str | int | bool]]) -> None:
    lines = [
        "-- MySQL data for Vietnamese ethnic groups",
        "-- Generated from dantoc/dantoc.docx",
        f"-- Generated at: {datetime.now().isoformat(timespec='seconds')}",
        f"-- Records: {len(groups)}",
        "",
        "SET NAMES utf8mb4;",
        "",
        "CREATE TABLE IF NOT EXISTS vn_ethnic_groups (",
        "  id TINYINT UNSIGNED NOT NULL COMMENT 'Order number from source list',",
        "  code CHAR(2) NOT NULL COMMENT 'Two-digit display code, e.g. 01, 02',",
        "  name VARCHAR(100) NOT NULL COMMENT 'Ethnic group name',",
        "  other_names TEXT NULL COMMENT 'Other names from source document',",
        "  subgroups TEXT NULL COMMENT 'Small groups / local groups from source document',",
        "  residence_area TEXT NULL COMMENT 'Residence areas from source document',",
        "  is_majority TINYINT(1) NOT NULL DEFAULT 0 COMMENT '1 for Kinh/Viet, 0 for others',",
        "  sort_order TINYINT UNSIGNED NOT NULL,",
        "  PRIMARY KEY (id),",
        "  UNIQUE KEY uq_vn_ethnic_groups_code (code),",
        "  KEY idx_vn_ethnic_groups_name (name),",
        "  KEY idx_vn_ethnic_groups_sort_order (sort_order)",
        ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;",
        "",
        "DELETE FROM vn_ethnic_groups;",
        "",
        "INSERT INTO vn_ethnic_groups (id, code, name, other_names, subgroups, residence_area, is_majority, sort_order) VALUES",
    ]

    values = []
    for group in groups:
        values.append(
            "  ("
            + ", ".join(
                [
                    str(group["id"]),
                    sql_str(str(group["code"])),
                    sql_str(str(group["name"])),
                    sql_str(str(group["other_names"])),
                    sql_str(str(group["subgroups"])),
                    sql_str(str(group["residence_area"])),
                    "1" if group["is_majority"] else "0",
                    str(group["sort_order"]),
                ]
            )
            + ")"
        )
    lines.append(",\n".join(values) + ";")
    lines.extend(
        [
            "",
            "-- Example: ethnic group combobox",
            "-- SELECT id, name FROM vn_ethnic_groups ORDER BY sort_order;",
            "",
            "-- Example: search by name / other names",
            "-- SELECT id, name, other_names FROM vn_ethnic_groups",
            "-- WHERE name LIKE CONCAT('%', ?, '%') OR other_names LIKE CONCAT('%', ?, '%')",
            "-- ORDER BY sort_order;",
            "",
        ]
    )

    OUTPUT_SQL.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    groups = extract_groups()
    if len(groups) != 54:
        raise RuntimeError(f"Expected 54 ethnic groups, got {len(groups)}")

    write_sql(groups)
    write_json(groups)
    print(f"created_sql={OUTPUT_SQL}")
    print(f"created_json={OUTPUT_JSON}")
    print(f"records={len(groups)}")


if __name__ == "__main__":
    main()
